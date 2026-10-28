import importlib
import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import identity
from database import connect, postgres_sql


def spreadsheet(rows):
    workbook = Workbook()
    workbook.active.append(identity.HEADERS)
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


VALID_ID = "0012345679"
SECOND_ID = "0084575948"


class ValidationTests(unittest.TestCase):
    def test_national_id(self):
        self.assertTrue(identity.valid_national_id(VALID_ID))
        for value in ("0000000000", "123", "0012345678", "۰۰۱۲۳۴۵۶۷۹"):
            self.assertFalse(identity.valid_national_id(value))
        self.assertEqual(identity.normalize("۰۰۱۲۳۴۵۶۷۹"), VALID_ID)

    def test_template_and_text(self):
        content = identity.workbook_bytes([(VALID_ID, "کاربر", "فعال", "=private")])
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.active["A2"].data_type, "s")
        self.assertEqual(workbook.active["A2"].number_format, "@")
        self.assertEqual(workbook.active["D2"].data_type, "s")
        self.assertEqual(identity.parse_workbook(content)[0][0], VALID_ID)

    def test_duplicate_numeric_inactive_formula_and_empty(self):
        row = (VALID_ID, "کاربر", "فعال", "code")
        for rows in ([row, row], [(12345679, "کاربر", "فعال", "code")],
                     [(VALID_ID, "کاربر", "نامعتبر", "code")],
                     [(VALID_ID, "=1+1", "فعال", "code")], []):
            with self.subTest(rows=rows), self.assertRaises(identity.ValidationError):
                identity.parse_workbook(spreadsheet(rows))

    def test_bad_file_and_headers(self):
        with self.assertRaises(identity.ValidationError):
            identity.parse_workbook(b"not xlsx")
        workbook = Workbook()
        workbook.active.append(["wrong"])
        output = BytesIO()
        workbook.save(output)
        with self.assertRaises(identity.ValidationError):
            identity.parse_workbook(output.getvalue())

    def test_sql_compatibility(self):
        self.assertEqual(postgres_sql("SELECT 1 WHERE parent_id IS ?"),
                         "SELECT 1 WHERE parent_id IS NOT DISTINCT FROM %s")
        self.assertIn("ON CONFLICT DO NOTHING", postgres_sql("INSERT OR IGNORE INTO users VALUES(?)"))


class DiscountValidationTests(unittest.TestCase):
    def test_discount_template_and_validation(self):
        content = identity.workbook_bytes([(VALID_ID, "section-code")], identity.DISCOUNT_HEADERS)
        self.assertEqual(identity.parse_workbook(content, discounts_only=True), [(2, VALID_ID, "section-code")])
        for rows in ([(VALID_ID, "one"), (VALID_ID, "two")], [(12345679, "code")], [(VALID_ID, "=1+1")]):
            workbook = Workbook()
            workbook.active.append(identity.DISCOUNT_HEADERS)
            for row in rows:
                workbook.active.append(row)
            output = BytesIO()
            workbook.save(output)
            with self.subTest(rows=rows), self.assertRaises(identity.ValidationError):
                identity.parse_workbook(output.getvalue(), discounts_only=True)

    def test_discount_section_title(self):
        self.assertEqual(identity.section_title("  تخفیف کتاب  "), "تخفیف کتاب")
        for title in ("", " " * 5, "title\nline", "x" * 81):
            with self.subTest(title=title), self.assertRaises(ValueError):
                identity.section_title(title)


class RoutingTests(unittest.TestCase):
    def setUp(self):
        self.bot = importlib.import_module("bot")
        self.bot.STATES.clear()

    def test_main_menu_inline_layout(self):
        for admin in (False, True):
            for message_id in (None, 7):
                with self.subTest(admin=admin, message_id=message_id), \
                     patch.object(self.bot, "child_sections", return_value=[(1, "راهنما", "")]), \
                     patch.object(self.bot, "get_setting", side_effect=lambda key, default: "0" if key == "feature_quiz" else default), \
                     patch.object(identity, "discount_sections", return_value=[(2, "تخفیف", 1)]), \
                     patch.object(self.bot, "is_admin", return_value=admin), \
                     patch.object(self.bot, "send_message") as send, \
                     patch.object(self.bot, "edit_message") as edit:
                    self.bot.show_menu(42, message_id)
                    target, unused = (edit, send) if message_id else (send, edit)
                    target.assert_called_once()
                    unused.assert_not_called()
                    markup = target.call_args.args[-1]
                    self.assertEqual(set(markup), {"inline_keyboard"})
                    rows = markup["inline_keyboard"]
                    self.assertTrue(all(len(row) == 2 for row in rows[:-1]))
                    self.assertIn(len(rows[-1]), (1, 2))
                    callbacks = [button["callback_data"] for row in rows for button in row]
                    self.assertIn("s1", callbacks)
                    self.assertIn("discount:2", callbacks)
                    self.assertIn("sched", callbacks)
                    self.assertNotIn("quiz", callbacks)
                    self.assertEqual("a:panel" in callbacks, admin)

    def test_start_requires_identity(self):
        with patch.object(self.bot, "has_access", return_value=False), \
             patch.object(self.bot, "send_message") as send, \
             patch.object(self.bot, "show_menu") as menu:
            self.bot.handle_message(42, "/start", "name")
            send.assert_called_once_with(42, "لطفاً کد ملی خود را وارد کنید")
            menu.assert_not_called()

    def test_callback_blocked(self):
        with patch.object(self.bot, "has_access", return_value=False), \
             patch.object(self.bot, "send_message"), \
             patch.object(self.bot, "answer_callback"), \
             patch.object(self.bot, "show_menu") as menu:
            self.bot.handle_callback(42, "menu", "cb", 1, "name")
            menu.assert_not_called()

    def test_admin_bypass(self):
        with patch.object(self.bot, "is_admin", return_value=True), \
             patch.object(identity, "lookup") as lookup:
            self.assertTrue(self.bot.has_access(42))
            lookup.assert_not_called()

    def test_login_success_and_denial(self):
        for allowed in (False, True):
            with self.subTest(allowed=allowed), \
                 patch.object(self.bot, "has_access", return_value=False), \
                 patch.object(identity, "bind", return_value=allowed), \
                 patch.object(self.bot, "register_user"), \
                 patch.object(self.bot, "show_menu") as menu, \
                 patch.object(self.bot, "send_message") as send:
                self.bot.handle_message(42, VALID_ID, "name")
                if allowed:
                    menu.assert_called_once()
                    send.assert_not_called()
                else:
                    menu.assert_not_called()
                    send.assert_called_once_with(42, "شما به این بات دسترسی ندارید")

    def test_missing_discount(self):
        with patch.object(self.bot, "has_access", return_value=True), \
             patch.object(self.bot, "rate_ok", return_value=True), \
             patch.object(self.bot, "answer_callback"), \
             patch.object(identity, "personal_discount", return_value=("تخفیف", "")), \
             patch.object(self.bot, "send_message") as send:
            self.bot.handle_callback(42, "discount", "cb", 1, "name")
            send.assert_called_once_with(42, "برای حساب شما کد تخفیف ثبت نشده است")

    def test_discount_scoped_to_sender(self):
        with patch.object(self.bot, "has_access", return_value=True), \
             patch.object(self.bot, "rate_ok", return_value=True), \
             patch.object(self.bot, "answer_callback"), \
             patch.object(identity, "personal_discount", return_value=("تخفیف", "private-code")) as lookup, \
             patch.object(self.bot, "send_message") as send:
            self.bot.handle_callback(42, "discount", "cb", 1, "name")
            lookup.assert_called_once_with(42, None)
            send.assert_called_once_with(42, "تخفیف\nprivate-code")

    def test_management_rejects_non_admin(self):
        with patch.object(self.bot, "is_admin", return_value=False), \
             patch.object(identity, "export_users") as export:
            self.bot.handle_identity_callback(42, "a:identity:export")
            export.assert_not_called()

    def test_selected_discount_and_inactive_section(self):
        with patch.object(self.bot, "has_access", return_value=True), \
             patch.object(self.bot, "rate_ok", return_value=True), \
             patch.object(self.bot, "answer_callback"), \
             patch.object(identity, "personal_discount", return_value=None) as lookup, \
             patch.object(self.bot, "send_message") as send:
            self.bot.handle_callback(42, "discount:7", "cb", 1, "name")
            lookup.assert_called_once_with(42, 7)
            send.assert_called_once_with(42, "این بخش تخفیف در دسترس نیست.")

    def test_discount_management_requires_admin(self):
        with patch.object(self.bot, "is_admin", return_value=False), \
             patch.object(identity, "create_discount_section") as create, \
             patch.object(self.bot, "send_message") as send:
            self.bot.handle_discounts_callback(42, "a:discounts:add")
            create.assert_not_called()
            send.assert_not_called()
            self.assertNotIn(42, self.bot.STATES)

    def test_discount_create_and_rename_states(self):
        for state in ("discount_add", "discount_rename"):
            with self.subTest(state=state), \
                 patch.object(self.bot, "is_admin", return_value=True), \
                 patch.object(identity, "create_discount_section", return_value=7) as create, \
                 patch.object(identity, "rename_discount_section", return_value=True) as rename, \
                 patch.object(self.bot, "discounts_admin_panel") as panel:
                self.bot.STATES[42] = [state, 7]
                self.assertTrue(self.bot.handle_state_message(42, "تخفیف کتاب", "name"))
                panel.assert_called_once_with(42, 7)
                self.assertNotIn(42, self.bot.STATES)
                if state == "discount_add":
                    create.assert_called_once_with("تخفیف کتاب")
                else:
                    rename.assert_called_once_with(7, "تخفیف کتاب")


@unittest.skipUnless(os.environ.get("TEST_POSTGRES") == "1", "Requires dedicated test PostgreSQL")
class PostgreSQLTests(unittest.TestCase):
    def setUp(self):
        identity.initialize()
        with connect() as connection:
            connection.execute("TRUNCATE user_discounts, discount_sections, authorized_users, identity_imports RESTART IDENTITY")
        identity.initialize()

    def test_multiple_sections_and_legacy_preserved(self):
        identity.import_users(spreadsheet([(VALID_ID, "کاربر", "فعال", "legacy")]), 1)
        self.assertTrue(identity.bind(42, VALID_ID))
        section_id = identity.create_discount_section("کتاب")
        content = identity.workbook_bytes([(VALID_ID, "book-code")], identity.DISCOUNT_HEADERS)
        self.assertEqual(identity.import_discounts(content, section_id), (1, 0))
        self.assertEqual(identity.import_discounts(content, section_id), (0, 1))
        self.assertEqual(identity.personal_discount(42, section_id), ("کتاب", "book-code"))
        self.assertEqual(identity.personal_discount(42, None)[1], "legacy")
        self.assertIsNone(identity.personal_discount(43, section_id))
        identity.rename_discount_section(section_id, "تخفیف کتاب")
        self.assertEqual(identity.personal_discount(42, section_id), ("تخفیف کتاب", "book-code"))
        identity.toggle_discount_section(section_id)
        self.assertIsNone(identity.personal_discount(42, section_id))
        self.assertNotIn(section_id, [row[0] for row in identity.discount_sections()])
        self.assertIn(section_id, [row[0] for row in identity.discount_sections(include_inactive=True)])
        identity.toggle_discount_section(section_id)
        exported = identity.export_discounts(section_id)
        self.assertEqual(identity.parse_workbook(exported, discounts_only=True)[0][2], "book-code")

    def test_unknown_user_rolls_back_discount_import(self):
        identity.import_users(spreadsheet([(VALID_ID, "کاربر", "فعال", "legacy")]), 1)
        section_id = identity.create_discount_section("کتاب")
        content = identity.workbook_bytes([(VALID_ID, "book-code"), (SECOND_ID, "unknown")], identity.DISCOUNT_HEADERS)
        with self.assertRaises(identity.ValidationError):
            identity.import_discounts(content, section_id)
        with connect() as connection:
            self.assertEqual(connection.execute("SELECT count(*) FROM user_discounts").fetchone()[0], 0)

    def test_import_bind_restart_update_revoke(self):
        content = spreadsheet([(VALID_ID, "کاربر", "فعال", "first")])
        self.assertEqual(identity.import_users(content, 1), (1, 0))
        self.assertFalse(identity.bind(42, "0012345678"))
        self.assertTrue(identity.bind(42, VALID_ID))
        self.assertFalse(identity.bind(43, VALID_ID))
        importlib.reload(identity)
        self.assertEqual(identity.lookup(42), (VALID_ID, "first"))
        self.assertEqual(identity.import_users(spreadsheet([(VALID_ID, "کاربر", "فعال", "second")]), 1), (0, 1))
        self.assertEqual(identity.lookup(42)[1], "second")
        identity.import_users(spreadsheet([(VALID_ID, "کاربر", "غیرفعال", "")]), 1)
        self.assertIsNone(identity.lookup(42))
        self.assertFalse(identity.bind(43, VALID_ID))

    def test_atomic_validation(self):
        rows = [(VALID_ID, "کاربر", "فعال", "code"), ("bad", "نام", "فعال", "")]
        with self.assertRaises(identity.ValidationError):
            identity.import_users(spreadsheet(rows), 1)
        with connect() as connection:
            self.assertEqual(connection.execute("SELECT count(*) FROM authorized_users").fetchone()[0], 0)

    def test_one_account_cannot_claim_two_records(self):
        identity.import_users(spreadsheet([
            (VALID_ID, "اول", "فعال", "one"),
            (SECOND_ID, "دوم", "فعال", "two"),
        ]), 1)
        self.assertTrue(identity.bind(42, VALID_ID))
        self.assertFalse(identity.bind(42, SECOND_ID))
        self.assertEqual(identity.lookup(42), (VALID_ID, "one"))

    def test_atomic_database_failure(self):
        with connect() as connection:
            connection.execute("ALTER TABLE authorized_users ADD CONSTRAINT reject_test CHECK (discount_code <> 'reject')")
        try:
            rows = [(VALID_ID, "کاربر", "فعال", "code"), (SECOND_ID, "کاربر دوم", "فعال", "reject")]
            with self.assertRaises(Exception):
                identity.import_users(spreadsheet(rows), 1)
            with connect() as connection:
                self.assertEqual(connection.execute("SELECT count(*) FROM authorized_users").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT count(*) FROM identity_imports").fetchone()[0], 0)
        finally:
            with connect() as connection:
                connection.execute("ALTER TABLE authorized_users DROP CONSTRAINT reject_test")

    def test_concurrent_binding_and_import(self):
        content = spreadsheet([(VALID_ID, "کاربر", "فعال", "code")])
        with ThreadPoolExecutor(max_workers=2) as executor:
            counts = list(executor.map(lambda admin: identity.import_users(content, admin), (1, 2)))
        self.assertEqual(sorted(counts), [(0, 1), (1, 0)])
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda account: identity.bind(account, VALID_ID), (42, 43)))
        self.assertEqual(sum(results), 1)


if __name__ == "__main__":
    unittest.main()
