from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from psycopg.errors import UniqueViolation

from database import connect

HEADERS = ("کد ملی", "نام و نام خانوادگی", "وضعیت دسترسی", "کد تخفیف اختصاصی")
DISCOUNT_HEADERS = ("کد ملی", "کد تخفیف اختصاصی")
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 100000


def initialize():
    with connect() as connection:
        connection.execute("""CREATE TABLE IF NOT EXISTS authorized_users (
            national_id VARCHAR(10) PRIMARY KEY CHECK (national_id ~ '^[0-9]{10}$'),
            full_name TEXT NOT NULL,
            active BOOLEAN NOT NULL,
            discount_code TEXT NOT NULL DEFAULT '',
            bale_id BIGINT UNIQUE,
            CHECK (bale_id IS NULL OR bale_id > 0))""")
        connection.execute("""CREATE TABLE IF NOT EXISTS identity_imports (
            id BIGSERIAL PRIMARY KEY, admin_id BIGINT NOT NULL,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            inserted INTEGER NOT NULL, updated INTEGER NOT NULL)""")
        connection.execute("""CREATE TABLE IF NOT EXISTS discount_sections (
            id BIGSERIAL PRIMARY KEY, title TEXT NOT NULL,
            legacy_key TEXT UNIQUE, active BOOLEAN NOT NULL DEFAULT TRUE)""")
        connection.execute("""INSERT INTO discount_sections(title, legacy_key)
            VALUES ('دریافت کد تخفیف', 'default') ON CONFLICT DO NOTHING""")
        connection.execute("""CREATE TABLE IF NOT EXISTS user_discounts (
            section_id BIGINT NOT NULL REFERENCES discount_sections(id),
            national_id VARCHAR(10) NOT NULL REFERENCES authorized_users(national_id),
            code TEXT NOT NULL, PRIMARY KEY(section_id, national_id))""")


def normalize(value):
    return value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")).strip()


def valid_national_id(value):
    if len(value) != 10 or not value.isascii() or not value.isdigit() or len(set(value)) == 1:
        return False
    remainder = sum(int(digit) * (10 - index) for index, digit in enumerate(value[:9])) % 11
    return int(value[-1]) == (remainder if remainder < 2 else 11 - remainder)


def lookup(bale_id):
    with connect() as connection:
        return connection.execute(
            "SELECT national_id, discount_code FROM authorized_users WHERE bale_id=%s AND active",
            (bale_id,),
        ).fetchone()


def bind(bale_id, national_id):
    if not valid_national_id(national_id):
        return False
    try:
        with connect() as connection:
            return bool(connection.execute("""UPDATE authorized_users SET bale_id=%s
                WHERE national_id=%s AND active AND (bale_id IS NULL OR bale_id=%s)
                RETURNING national_id""", (bale_id, national_id, bale_id)).fetchone())
    except UniqueViolation:
        return False


def workbook_bytes(rows=(), headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "کاربران"
    sheet.sheet_view.rightToLeft = True
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    for index, width in enumerate((20, 35, 22, 30), 1):
        sheet.column_dimensions[chr(64 + index)].width = width
        sheet.column_dimensions[chr(64 + index)].number_format = "@"
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
    for row in sheet.iter_rows(min_row=2, max_row=max(1001, sheet.max_row), max_col=len(headers)):
        for cell in row:
            cell.number_format = "@"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_users():
    with connect() as connection:
        rows = connection.execute("""SELECT national_id, full_name,
            CASE WHEN active THEN 'فعال' ELSE 'غیرفعال' END, discount_code
            FROM authorized_users ORDER BY national_id""").fetchall()
    return workbook_bytes(rows)


class ValidationError(ValueError):
    def __init__(self, errors, count):
        self.errors = errors
        self.count = count
        super().__init__("\n".join(errors[:15]))


def parse_workbook(content, discounts_only=False):
    if len(content) > MAX_FILE_SIZE:
        raise ValidationError(["حجم فایل بیش از ۵ مگابایت است."], 1)
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError(["حجم بازشده فایل بیش از حد مجاز است."], 1)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except Exception:
        raise ValidationError(["فایل اکسل معتبر با پسوند xlsx ارسال کنید."], 1) from None
    try:
        if len(workbook.worksheets) != 1:
            raise ValidationError(["فایل باید دقیقاً یک برگه داشته باشد."], 1)
        sheet = workbook.active
        rows = sheet.iter_rows()
        header = next(rows, ())
        expected_headers = DISCOUNT_HEADERS if discounts_only else HEADERS
        if tuple(cell.value for cell in header) != expected_headers:
            raise ValidationError(["ستون‌ها باید دقیقاً مطابق قالب باشند."], 1)
        records, seen, errors = [], set(), []
        error_count = 0
        for number, cells in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise ValidationError(["تعداد ردیف‌ها بیش از حد مجاز است."], 1)
            values = [cell.value for cell in cells]
            if all(value is None for value in values):
                continue
            reason = None
            if len(values) != len(expected_headers) or any(cell.data_type == "f" for cell in cells):
                reason = "تعداد ستون نادرست یا فرمول غیرمجاز"
            elif not isinstance(values[0], str):
                reason = "کد ملی باید از نوع متن باشد؛ صفرهای ابتدایی را بررسی کنید"
            else:
                national_id = normalize(values[0])
                if not valid_national_id(national_id):
                    reason = "کد ملی نامعتبر"
                elif national_id in seen:
                    reason = "کد ملی تکراری"
                elif discounts_only:
                    if values[1] is not None and not isinstance(values[1], str):
                        reason = "کد تخفیف باید متن باشد"
                    else:
                        seen.add(national_id)
                        records.append((number, national_id, values[1] or ""))
                elif not isinstance(values[1], str) or not values[1].strip():
                    reason = "نام و نام خانوادگی الزامی است"
                elif values[2] not in ("فعال", "غیرفعال"):
                    reason = "وضعیت باید فعال یا غیرفعال باشد"
                elif values[3] is not None and not isinstance(values[3], str):
                    reason = "کد تخفیف باید متن باشد"
                else:
                    seen.add(national_id)
                    records.append((national_id, values[1].strip(), values[2] == "فعال", values[3] or ""))
            if reason:
                error_count += 1
                if len(errors) < 15:
                    errors.append(f"ردیف {number}: {reason}")
        if error_count:
            raise ValidationError(errors, error_count)
        if not records:
            raise ValidationError(["فایل فاقد ردیف اطلاعات است."], 1)
        return records
    finally:
        workbook.close()


def import_users(content, admin_id):
    records = parse_workbook(content)
    with connect() as connection:
        connection.execute("LOCK TABLE authorized_users IN SHARE ROW EXCLUSIVE MODE")
        identifiers = [record[0] for record in records]
        updated = connection.execute(
            "SELECT count(*) FROM authorized_users WHERE national_id = ANY(%s)", (identifiers,)
        ).fetchone()[0]
        with connection.cursor() as cursor:
            cursor.executemany("""INSERT INTO authorized_users(national_id, full_name, active, discount_code)
                VALUES (%s,%s,%s,%s) ON CONFLICT(national_id) DO UPDATE SET
                full_name=excluded.full_name, active=excluded.active,
                discount_code=excluded.discount_code""", records)
        inserted = len(records) - updated
        connection.execute("INSERT INTO identity_imports(admin_id, inserted, updated) VALUES (%s,%s,%s)",
                           (admin_id, inserted, updated))
    return inserted, updated


def discount_sections(include_inactive=False):
    with connect() as connection:
        return connection.execute("""SELECT id, title, active FROM discount_sections
            WHERE active OR %s ORDER BY id""", (include_inactive,)).fetchall()


def section_title(title):
    title = title.strip()
    if not title or len(title) > 80 or any(ord(character) < 32 for character in title):
        raise ValueError("نام بخش باید بین ۱ تا ۸۰ نویسه و در یک خط باشد.")
    return title


def create_discount_section(title):
    with connect() as connection:
        return connection.execute("INSERT INTO discount_sections(title) VALUES (%s) RETURNING id",
                                  (section_title(title),)).fetchone()[0]


def rename_discount_section(section_id, title):
    with connect() as connection:
        return bool(connection.execute("UPDATE discount_sections SET title=%s WHERE id=%s RETURNING id",
                                       (section_title(title), section_id)).fetchone())


def toggle_discount_section(section_id):
    with connect() as connection:
        connection.execute("UPDATE discount_sections SET active=NOT active WHERE id=%s", (section_id,))


def personal_discount(bale_id, section_id):
    with connect() as connection:
        return connection.execute("""SELECT sections.title,
            CASE WHEN sections.legacy_key='default' THEN users.discount_code ELSE codes.code END
            FROM authorized_users users CROSS JOIN discount_sections sections
            LEFT JOIN user_discounts codes ON codes.section_id=sections.id AND codes.national_id=users.national_id
            WHERE users.bale_id=%s AND users.active AND sections.active
            AND (sections.id=%s OR (%s::bigint IS NULL AND sections.legacy_key='default'))""",
            (bale_id, section_id, section_id)).fetchone()


def export_discounts(section_id):
    with connect() as connection:
        section = connection.execute("SELECT legacy_key FROM discount_sections WHERE id=%s", (section_id,)).fetchone()
        if not section:
            raise ValueError("بخش یافت نشد.")
        rows = connection.execute("""SELECT users.national_id,
            CASE WHEN %s THEN users.discount_code ELSE COALESCE(codes.code, '') END
            FROM authorized_users users LEFT JOIN user_discounts codes
            ON codes.national_id=users.national_id AND codes.section_id=%s ORDER BY users.national_id""",
            (section[0] == 'default', section_id)).fetchall()
    return workbook_bytes(rows, DISCOUNT_HEADERS)


def import_discounts(content, section_id):
    numbered_records = parse_workbook(content, discounts_only=True)
    records = [(national_id, code) for number, national_id, code in numbered_records]
    with connect() as connection:
        connection.execute("LOCK TABLE authorized_users IN SHARE ROW EXCLUSIVE MODE")
        section = connection.execute("SELECT legacy_key FROM discount_sections WHERE id=%s FOR UPDATE",
                                     (section_id,)).fetchone()
        if not section:
            raise ValueError("بخش یافت نشد.")
        known = {row[0] for row in connection.execute(
            "SELECT national_id FROM authorized_users WHERE national_id=ANY(%s)",
            ([record[0] for record in records],)).fetchall()}
        missing = [number for number, national_id, code in numbered_records if national_id not in known]
        if missing:
            raise ValidationError([f"ردیف {index}: کاربر در فهرست احراز هویت ثبت نشده است"
                                   for index in missing[:15]], len(missing))
        updated = len(records) if section[0] == 'default' else connection.execute(
            "SELECT count(*) FROM user_discounts WHERE section_id=%s AND national_id=ANY(%s)",
            (section_id, [record[0] for record in records])).fetchone()[0]
        with connection.cursor() as cursor:
            if section[0] == 'default':
                cursor.executemany("UPDATE authorized_users SET discount_code=%s WHERE national_id=%s",
                                   [(code, national_id) for national_id, code in records])
            else:
                cursor.executemany("""INSERT INTO user_discounts(section_id, national_id, code)
                    VALUES (%s,%s,%s) ON CONFLICT(section_id, national_id) DO UPDATE SET code=excluded.code""",
                    [(section_id, national_id, code) for national_id, code in records])
    return len(records) - updated, updated
